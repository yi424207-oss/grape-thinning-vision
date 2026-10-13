"""Visual course paper; figures are tied to saved real files by an evidence index."""
from paper_layout import *
from docx.enum.text import WD_BREAK
berry=json.loads((R/'reports/berry_aux_evaluation.json').read_text())
flower=json.loads((R/'reports/flower_evaluation.json').read_text())
young=json.loads((R/'reports/young_evaluation.json').read_text())
base=json.loads((R/'reports/young_baseline_evaluation.json').read_text())
review=json.loads((R/'reports/review_statistics.json').read_text())
cpu=json.loads((R/'reports/cpu_benchmark.json').read_text())
import yaml
configs={task:yaml.safe_load((R/'runs'/task/'args.yaml').read_text()) for task in ['berry_aux','flower','young']}
def page(title):
 h(title);D.paragraphs[-1].paragraph_format.page_break_before=True
def note(text):
 a=D.add_paragraph(text);a.paragraph_format.first_line_indent=Pt(0)
 a.paragraph_format.line_spacing=Pt(15)
 for r in a.runs:r.font.size=Pt(9)
 md.append(text)
def pair(files,captions,width=2.8):
 a=D.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.line_spacing=1.0;a.paragraph_format.keep_with_next=True
 for i,file in enumerate(files):
  if i:a.add_run('  ')
  a.add_run().add_picture(str(O/'figures'/file),width=Inches(width))
 cap=D.add_paragraph('；'.join(captions),'Caption');cap.alignment=WD_ALIGN_PARAGRAPH.CENTER
 md.extend(f'![{c}](figures/{f})' for f,c in zip(files,captions))
# Cover
for text,style,before,after in [('生物生产机器人设计','Title',65,20),('基于 YOLO 与相对尺度判别的\n葡萄可见果粒检测与尺度筛选算法原型设计','Subtitle',0,38)]:
 a=D.add_paragraph(text,style);a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_before=Pt(before);a.paragraph_format.space_after=Pt(after)
for text in ['课程设计选题  3 葡萄疏果识别图像处理算法','姓名  ____________________','学号  ____________________','专业班级  ____________________','指导教师  ____________________','完成日期  2026年10月']:
 a=D.add_paragraph(text);a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Pt(0)
page('摘要与设计目标')
p('针对葡萄疏果中的可见果粒计数与小果筛选需求，本文设计基于YOLOv8n的检测与相对尺度筛选原型。花序分支输出整花序与小穗位置，果粒分支结合单穗ROI、尺度规则和可见保留上限生成编号图及CSV、JSON记录，供人工复查。')
p('主实验中，辅助果粒120个测试图块的F1为0.7467，整花序和小穗在79张测试图上的F1分别为0.8865和0.5626；补充三个训练种子后，辅助果粒F1为0.7493±0.0015。整花序定位较稳定，小穗及稠密果粒的漏检仍是主要误差来源。幼果适配在2个测试区域上完成流程试验，F1为0.8800、MAE为3.00粒/区域。幼果标签未经专家复核，桌面交互尚未验收，当前交付依据命令行流程与可复查输出评价。')
p('关键词：阳光玫瑰；疏果；YOLOv8；可见果粒计数；相对尺度')
h('1 研究背景与任务要求')
p('传统颜色和形状方法容易把绿色叶片、果梗与幼果混淆。Mask-GK采用实例掩膜高斯核进行果粒分割计数[1]；MVDNet结合计数与果穗分割研究疏果期计数[2]。本文选择具有现成预训练权重和检测工具链的YOLOv8n[3]，将课程设计工作集中在数据处理、检测与筛选规则的衔接和输出核查。CPU训练条件及现有框标签使这一方案便于实施；未开展检测器横向比较，模型选型依据是实现条件，而非已证明其优于实例分割。')
p('疏果效果受品种、生育期及环境条件影响[9]。公开精品果指标中，马陆阳光玫瑰为45—50粒每串[4]，对应整穗生产目标。本文处理的是二维照片中可见对象的相对大小，保留上限由操作者根据ROI设定。网页栽培记录[5]用于图像来源与任务背景，尺度规则的农艺适用性在第8节讨论。')
table('表1 课程要求与实现',['要求','实现及边界'],[['花或小果图像','花序与幼果分别训练，较大果粒只作辅助'],['小果与去除阈值','相对尺度默认0.70，可调工程参数'],['留下数量','可见数减建议去除数，不含遮挡果粒'],['设计程序','CLI、桌面源码、模型及实际导出样例']])
page('2 数据样本与划分')
p('Du花序包采用CC BY 4.0许可[6]。在567份命名条目中发现44份图片与标签完全一致的重复，清洗后保留523张原图，并排除原包已有增强图，避免相同内容跨划分。GBISC包含巨峰、阳光玫瑰和夏黑各50张原图[1]。辅助模型混合使用三品种，以利用已有果粒框并覆盖不同外观；品种效果单独列于表2。每品种按连续5张原图为一组，以种子20261010将10组分为6训练、2验证、2测试，每品种原图数为30/10/10。划分先于四象限裁切，同源图块留在同一侧，控制裁切造成的已知泄漏。')
fig('berry_samples.png','图1 三品种辅助果粒训练图块与标签 GBISC CC BY 4.0',4.3)
table('表2 三品种原图划分与固定阈值测试',['品种','训练 验证 测试原图','测试F1'],[[name,'30 10 10',f"{review['cultivar_results'][key]['f1']:.4f}"] for name,key in [('巨峰','Kyoho'),('阳光玫瑰','ShineMuscat'),('夏黑','SummerBlack')]])
page('2.1 花序训练样本')
p('花序模型包含整花序与小穗两类。下图保留实际文件名，分别显示不同穗型、背景与视角的标注。清洗后按拍摄序号分组，原包的重复文件和已有旋转增强不参与独立原图计数。')
fig('flower_samples.png','图2 花序训练原图与标签 Du CC BY 4.0',4.3)
page('2.2 预处理与幼果适配')
table('表3 实际划分与标签性质',['数据','训练','验证','测试','标签'],[['辅助果粒','360图块','120图块','120图块','原实例框'],['花序','369图','75图','79图','花序及小穗框'],['幼果','37图','1区域','2区域','训练伪标签 评测视觉框']])
p('辅助果粒图块对应90/30/30张原图。裁切后目标可见框面积须达到原框的65%，最长边缩至768像素。花序按拍摄编号整除10分组，再以种子20261010打乱组号划分；逐组数量见reports/split_group_distribution.json。编号组提供可核查的划分单位，植株及果园身份缺失的影响集中见第8节。')
p('幼果试验用于检验辅助模型向疏果期图像适配后，能否贯通检测、尺度筛选与导出流程。按品种配文及阶段筛出37张网络照片[7]，由辅助模型生成1571个训练伪标签。同网页处理前后图均进入训练侧；验证与测试另取自两个网站[5,8]，评测框经AI视觉辅助逐框检查。采用弱标注是现有样本和标注条件下的探索性安排，因此幼果数据单列讨论，不用于总体性能排序。公开交付的数据范围及复现条件见第8节。')
pair(['berry_aux_train_batch0.jpg','flower_train_batch0.jpg'],['图3a 辅助果粒实际训练批次','图3b 花序实际训练批次'],2.75)
note('图3由训练器实际保存，展示缩放与拼接增强后的输入。它用于核对训练过程，独立测试结果见第5节。')
page('3 算法结构与设计程序')
fig('system_design.png','图4 检测 筛选及结果导出的功能关系',6.0)
p('输入图像先由用户框选单穗。大区域使用1024像素窗口和768像素步长，坐标回到原图后以NMS消除重复框。模型输入固定填充到384像素，避免单图与批量填充方式不同造成计数差异。花序分支用于定位整花序、统计小穗及保存框坐标，不自动决定疏花或疏小穗；其计数不换算成果粒数，也不与果粒去除动作相加。果粒分支才继续执行尺度筛选。')
p('绝对像素大小随拍摄距离和成像比例变化，因而以当前ROI内检测框的中位尺度作归一化。第i个框宽高为wᵢ、hᵢ，投影尺度dᵢ=√(wᵢhᵢ)，相对尺度rᵢ=dᵢ/median(d)。当rᵢ<τ时建议去除。τ=0.70为可调的默认操作点，表7检查0.50至0.90范围内的规则响应；该值未经农艺目标优化。rᵢ表示二维相对投影尺度，尺度基准对选择性漏检的敏感性见7.1节。')
p('保留上限K用于把操作者给定的可见目标转化为数量约束。在大小筛选后，从中心距离除以平均尺度最小的一对中优先去除较小者，迭代至保留数不超过K。N_total=N_keep+N_remove分别对应检测总数、保留数和建议去除数。逐粒编号、位置、尺度与原因共同保存，便于检查规则是否按预期执行。')
rows=[['初始权重','COCO YOLOv8n','COCO YOLOv8n','辅助 best.pt']]
for label,key in [('轮次','epochs'),('输入边长 px','imgsz'),('批量','batch'),('设备','device'),('优化器','optimizer'),('初始学习率','lr0'),('冻结层数','freeze'),('随机种子','seed')]:
 rows.append([label]+[str(configs[t][key]) for t in ['berry_aux','flower','young']])
page('3.1 原始模型训练配置')
table('表4 各模型独立训练配置',['参数','辅助果粒','花序与小穗','幼果适配'],rows)
note('原训练日志记录Python 3.12.14、PyTorch 2.14.1+cpu、Ultralytics 8.4.175。依赖已固定在requirements.txt，补充的完整CPU依赖锁定文件为requirements-lock-cpu.txt。原主实验各只有一个随机种子；本轮追加每模型三个种子复验，配置和结果单独报告。')
for task,title,num,detail in [
 ('berry_aux','4 辅助果粒模型训练结果',5,'图5中训练损失总体下降，验证指标逐渐趋稳。图6的测试FN为3393，高于FP的827，主要困难表现为漏检；其对计数误差的影响见第5节。'),
 ('flower','4.1 花序与小穗模型训练结果',7,'整花序测试TP为82、FP为8、FN为13，定位较稳定；小穗为TP=879、FP=541、FN=826，误检与漏检同时存在。小穗局部结构相似且边界难分，可能增加定位难度。spikelet在数据中指小穗。'),
 ('young','4.2 幼果适配流程试验',9,'图9显示验证指标存在明显波动。验证只含1个区域，少量框匹配变化即可改变指标；因此曲线主要用于检查适配过程，幼果输出单列于5.1节。')]:
 page(title)
 fig(f'{task}_results.png',f'图{num} 实际results.csv对应的训练损失及验证指标 results.png',6.25)
 fig(f'{task}_cm_evidence.png',f'图{num+1} 同一置信度与匹配IoU下验证和测试的TP FP FN',6.1)
 p(detail)
 if task=='berry_aux':
  note('图5的results.png由保存的训练CSV生成，沿用原训练器AP口径：NMS IoU=0.70、最多300框，用于观察训练趋势。图6按保存预测重新匹配，IoU≥0.50，置信度由验证集选择并固定用于测试。TP为一对一正确匹配，FP为未匹配预测，FN为未检出标签；检测评价不定义背景真阴性。F1=2TP/(2TP+FP+FN)。原conf=0.001、匹配IoU=0.45验证矩阵仅作过程档案。后续训练图与TP/FP/FN图沿用本节说明，AP积分口径见7.2节。')
 else:
  note('训练曲线及TP/FP/FN的统计口径同第4节。')
page('5 独立测试结果与计数误差')
p('表5汇总辅助果粒、整花序与小穗的主实验结果，采用第4节定义的匹配和阈值选择方法。图11保留各分支过程对照，其中幼果面板仅供流程核查，不参与主结果比较。辅助果粒F1为0.7467且计数MAE较大，说明检测到部分果粒与准确查数仍有距离；整花序F1为0.8865，小穗F1为0.5626，整穗定位和局部小目标识别的难度差异明显。')
fig('test_confusion.png','图11 固定阈值测试的TP FP FN及统一精度F1',6.3)
rows=[]
for name,d,key in [('辅助果粒',berry,'berry'),('整花序',flower,'inflorescence'),('小穗',flower,'spikelet')]:
 q=d['classes'][key]['test'];rows.append([name,str(q['images']),f"{q['threshold']:.2f}",f"{q['precision']:.3f}",f"{q['recall']:.3f}",f"{q['f1']:.4f}",f"{q['count_mae']:.2f}"])
table('表5 固定阈值测试结果',['对象','图数','阈值','P','R','F1','MAE'],rows)
fig('test_count_scatter.png','图12 测试图标签数量与预测数量 虚线为理想一致关系',6.1)
note('表5的MAE单位依次为粒/图块、穗/图和小穗/图。辅助120图块对应30张原图，计数误差按图块计算。图12的幼果散点仅含2个区域，另见5.1节。')
page('5.1 验证PR曲线与幼果流程试验')
p('幼果主实验的2个测试区域共78个视觉辅助检查框，F1=0.8800、MAE=3.00粒/区域。这一分支评价流程是否可执行，不据此与表5的任务排序。表6进一步在同一置信度下检查适配前后的输出差异。')
fig('flower_BoxPR_curve.png','图13 花序最佳权重重跑验证PR曲线 整花序与小穗分开显示',3.4)
pair(['berry_aux_BoxPR_curve.png','young_BoxPR_curve.png'],['图14a 辅助最佳权重验证PR','图14b 幼果最佳权重验证PR'],1.95)
rows=[]
for c in review['young_common_threshold_comparison']:
 for key,label in [('baseline','辅助直接预测'),('adapted','弱标注适配')]:
  q=c[key];rows.append([label,f"{c['common_threshold']:.2f}",f"{q['f1']:.4f}",f"{q['count_mae']:.2f}"])
table('表6 同一幼果测试区域的同阈值对照',['模型','共同阈值','F1','MAE 粒每区域'],rows)
note('在0.40阈值下，适配模型MAE由5.00降至3.00粒/区域；0.50下则由4.50升至6.50。检测匹配与数量误差呈现不同响应，说明适配后的收益依赖工作阈值，而非所有指标同步提高。PR曲线对应验证集，表6对应同一组测试区域。')
page('6 程序输入与实际输出')
fig('berry_input_output.png','图15 阳光玫瑰示例的输入与果粒导出 GBISC CC BY 4.0',5.1)
p('果粒命令行入口为src/predict.py，绿色框为保留，红色框为建议去除。所有数量只覆盖检测到的可见对象。多穗照片应先框选单穗，否则统计的是整个输入区域。')
fig('flower_input_output.png','图16 花序示例的输入与检测导出 Du CC BY 4.0',5.1)
p('花序入口为src/flower_predict.py，整花序与小穗计数分别输出。本页来自实际导出的annotated.jpg，展示花序分支的功能。')
page('6.1 运行记录与结构化导出')
fig('cli_run_record.png','图17 修订后实际命令行记录 ASCII双连字符参数',6.1)
fig('csv_export_preview.png','图18 实际导出记录中的保留和建议去除样例',6.1)
p('导出目录同时包含annotated.jpg、berries.csv与result.json。CSV采用UTF-8 BOM；schema_version=2.0使用relative_scale表示第3节定义的尺度。JSON记录ROI、模型、阈值、耗时与逐粒决策。图17给出实际命令和计数，图18同时展示keep与remove记录，使编号图和结构化输出能够相互核对；字段说明见表12。')
p('桌面入口src/app.py提供打开照片、拖框选ROI、切换模型和参数输入，推理线程异步回传结果。当前环境无法建立图形服务连接，尚未完成窗口交互验收，窗口操作仍需在目标电脑验收。已实际验证的入口是命令行。')
page('7 阈值响应与失败样例')
fig('test_error_examples.png','图19 独立测试中计数误差最大的两个辅助果粒图块 青色标签 橙色预测',6.1)
p('图19按测试集绝对计数误差选出两个失败图块：标签/预测数量分别为123/42和166/93，误差为−81和−73粒。图中可见目标重叠及叶片遮挡，384像素输入也压缩了局部细节，可能削弱相邻果粒的可分辨性。现有记录尚未分离输入分辨率、置信度和NMS各自的作用，故不将具体漏检归因于某一步。尺度阈值作用于检测后的框集合，提高τ只能改变已有框的动作，不能恢复未检出的果粒。')
sens=json.loads((R/'reports/threshold_sensitivity.json').read_text())
table('表7 同一40粒演示ROI的小果比例敏感性',['比例τ','检测数','建议去除','保留'],[[f"{x['size_ratio']:.2f}",x['visible_count'],x['remove_count'],x['keep_count']] for x in sens])
p('同一40框集合中，τ从0.50升至0.90时建议去除数由3增至19，响应单调；0.70下去除15、保留25。该检查验证规则实现及数量守恒，不用于选择农艺最优阈值。原版完成6项单元测试及路径、空白图、缺失文件、越界ROI和导出一致性检查；修订后共8项测试通过。')
page('7.1 选择性漏检与规则敏感性')
fig('scale_missingness_sensitivity.png','图20 删除最小框的模拟对剩余对象建议的影响',6.1)
note('在保存的40框示例上模拟删除最小的10%、20%、30%检测框，中位尺度由106.70 px增至110.02、116.72、121.71 px；剩余对象建议分别改变0、1、2个。这是检测结果上的扰动分析，不是新检测性能实验。中位数归一化也会受选择性漏检影响；分位数或空间分层需另行验证，当前建议必须由人工复查。')
page('7.2 指标不确定性与AP评估')
rows=[]
for task,name,label in [('berry_aux','berry','辅助果粒'),('flower','inflorescence','整花序'),('flower','spikelet','小穗')]:
 q=review['models'][task]['classes'][name]['conditional_group_bootstrap'];a=q['f1_percentile_95'];b=q['mae_percentile_95']
 rows.append([label,q['groups'],f'{a[0]:.4f} 至 {a[1]:.4f}',f'{b[0]:.2f} 至 {b[1]:.2f}'])
table('表8 按编号组重采样的条件百分位区间',['对象','组数','F1 95%区间','MAE 95%区间'],rows)
note('重采样10000次，随机种子20261012；辅助按连续5张原图组，花序按编号整除10组。区间以标签和当前模型为条件，并假设编号组独立；来源缺少植株身份，不能据此证明场景独立。重采样不能替代多种子重训，也不覆盖标注偏差。幼果只有两个同来源区域，不给出具有总体解释的置信区间。')
ap=json.loads((R/'reports/review_test_AP.json').read_text())
rows=[]
for task in ['berry_aux','flower']:
 for name,q in ap[task]['class_AP'].items():rows.append([task+' '+name,f"{q['AP50']:.4f}",f"{q['AP50_95']:.4f}"])
rows.append(['幼果 流程试验',f"{young['test_AP']['metrics/mAP50(B)']:.4f}",f"{young['test_AP']['metrics/mAP50-95(B)']:.4f}"])
table('表9 独立测试AP与固定阈值指标分开报告',['对象','AP50','AP50至95'],rows)
note('三组AP均已以原主实验权重重新验证，并与原记录一致；本轮重复训练另列。AP50匹配IoU为0.50；AP50至95为IoU 0.50至0.95步长0.05的平均。AP评测候选置信度下限0.001、NMS IoU 0.50，不能把该下限称作固定工作阈值。')
page('7.3 三种随机种子重复训练')
seed_report=json.loads((R/'reports/seed_repeats.json').read_text())
assert len(seed_report['runs'])==9, 'Complete all nine seed runs before publishing this paper'
p('原主实验各使用一个种子，缺少对训练随机性影响的估计，因此在固定划分下追加20261010、20261012、20261013三个种子，每次25轮，共9次。辅助果粒与花序均使用同一COCO预训练权重；幼果从原主实验辅助best.pt开始，种子变化主要作用于后续训练过程。冻结层数、优化器、学习率与输入大小沿用原配置；本轮为8线程、workers=0、patience=25。')
fig('seed_repeat_results.png','图21 固定划分下各随机种子的F1与计数MAE',6.1)
rows=[]
for task,name,label,unit in [('berry_aux','berry','辅助果粒','粒每图块'),('flower','inflorescence','整花序','穗每图'),('flower','spikelet','小穗','小穗每图'),('young','berry','幼果试验','粒每区域')]:
 q=seed_report['summary'][task][name];f=q['f1'];m=q['count_mae']
 rows.append([label,f"{f['mean']:.4f} ± {f['sample_std']:.4f}",f"{m['mean']:.2f} ± {m['sample_std']:.2f}",unit])
table('表10 三次训练的均值和样本标准差',['对象','F1','计数MAE','MAE单位'],rows)
note('辅助果粒三次F1为0.7493±0.0015，波动较小；幼果为0.8652±0.0279，MAE为3.00±1.00粒/区域，当前小样本下波动更明显。每次最佳轮次与工作阈值均由验证集选择，完整逐次指标见reports/seed_repeats.json。标准差描述固定划分和标签条件下的训练波动，幼果测试区域数仍为2。本轮硬件及workers与原主实验不同，二者分别列示。')
ap_seed=[]
for task,label in [('berry_aux','辅助果粒'),('flower','花序两类平均'),('young','幼果试验')]:
 a=seed_report['test_AP_summary'][task]['metrics/mAP50(B)'];b=seed_report['test_AP_summary'][task]['metrics/mAP50-95(B)']
 ap_seed.append(f"{label} {a['mean']:.4f}±{a['sample_std']:.4f} / {b['mean']:.4f}±{b['sample_std']:.4f}")
note('三种子测试mAP50 / mAP50至95的均值±标准差：'+'；'.join(ap_seed)+'。花序值为两类平均，不是单独小穗的AP。')
page('7.4 运行测量和导出字段')
rows=[]
for q in cpu['tasks']:
 rows.append([q['model'],q['batch'],f"{q['mean_per_image_seconds']*1000:.2f}",f"{q['throughput_images_per_second']:.1f}",f"{q['process_peak_rss_mib']:.1f}"])
table('表11 当前CPU预热推理测量',['模型','批量','ms每图','图每秒','峰值RSS MiB'],rows)
note('硬件为'+cpu['cpu']+'，固定8线程、384像素方形填充。预热3次后测20次；同一已解码许可图重复输入，包含预处理、前向与NMS，不包含模型加载、磁盘读取、ROI切窗和结果导出。吞吐量不是交互界面的端到端FPS，峰值RSS为进程累计峰值，不是独立模型内存。原训练硬件为AMD EPYC 9V74，不能与当前测量混同。原始时间序列与标准差见reports/cpu_benchmark.json。')
table('表12 CSV及JSON逐粒字段',['字段','含义和取值'],[['id','当前结果内从1开始的编号 对应编号图 不保证跨图一致'],['box 或 x1 y1 x2 y2','JSON使用box数组 CSV使用四个坐标列 单位为原图像素'],['confidence','检测置信度 0至1 不表示农艺正确概率'],['relative_scale','sqrt(框面积)除以当前检测框尺度中位数 无量纲'],['action','keep 保留 或 remove 建议去除'],['reason','keep 未触发规则 relative_small 低于尺度阈值 crowding_target 因可见上限去除']])
note('只有设置visible_target时才使用crowding_target；花序分支仅输出定位与类别计数，不输出果粒去除动作。8项单元测试通过；新版CLI实测39个检测对象，建议去除12、保留27，字段与编号图可对应。')
page('8 结论与局限')
p('本文贯通了可见目标检测、尺度筛选、可见保留上限与结构化导出。整花序定位是当前较稳定的识别环节；小穗和稠密果粒仍有较大漏检，单纯调整尺度规则难以解决查数误差。幼果分支完成了适配与程序流程检查，但它是证据最薄弱的部分：训练仅37张、测试仅2个区域，标签未经双人或农艺专家复核。本文据此区分检测性能、规则响应和流程可执行性。')
p('评价范围还受数据和验收条件约束。编号分组控制了同源图块跨划分，但缺少植株与果园身份，无法建立充分的场景独立性；bootstrap区间以当前标签、模型和分组假设为条件。二维检测框不覆盖隐藏果粒、病虫害、朝向和穗形等农艺因素，当前筛选仍需人工复查，尚无田间疏除或机械剪切真值。后续应优先扩充跨果园、跨光照的幼果独立样本并完成标签复核，再比较输入分辨率、检测后处理或实例分割方案，最后进行田间疏除验证。')
table('表13 交付包可直接查找的材料',['目录','内容'],[['data/berry_yolo images labels','实际辅助果粒训练360 验证120 测试120图块'],['data/flower_yolo images labels','实际花序训练369 验证75 测试79原图'],['幼果模型与汇总记录','公开包不包含幼果原图及逐图标注 无法仅依公开包重训'],['runs 各模型目录','best last权重 results.csv results.png 矩阵 PR'],['docs figures','论文图及图像来源证据索引'],['src scripts tests reports','程序 复现脚本 单元测试 实际评测与日志']])
p('交付包保留辅助果粒与花序的许可图像、原模型及复验权重。幼果网络原图及逐图标注未获明确再分发许可，公开包仅提供权重和汇总证据；网页失效也可能影响来源恢复。命令行入口已实际运行，桌面入口因图形服务连接限制尚未完成窗口交互验收。解压后按使用说明安装依赖，运行许可示例，更换目录后使用scripts/relink_datasets.py更新数据路径。')
page('参考文献')
refs=[
'[1] YANG C, GENG T, PENG J, et al. Mask-GK An efficient method based on mask Gaussian kernel for segmentation and counting of grape berries in field[J]. Computers and Electronics in Agriculture, 2025, 234:110286. DOI:10.1016/j.compag.2025.110286.',
'[2] DU W, QIN W, CUI X, et al. A precise berry counting method for in-cluster grapes to guide berry thinning[J]. Frontiers in Plant Science, 2026,16:1739688. DOI:10.3389/fpls.2025.1739688. 本设计未取得其完整779张数据.',
'[3] ULTRALYTICS. Explore Ultralytics YOLOv8[EB/OL]. https://docs.ultralytics.com/models/yolov8. 访问日期2026-10-10.',
'[4] 上海市人民政府新闻办公室. 精品果按串销售每串售价100元 马陆葡萄品牌价值再创新高[EB/OL]. https://www.shio.gov.cn/TrueCMS/shxwbgs/ywts/content/94f32927-3c25-4f87-82c2-1cf6438d6cc2.htm. 访问日期2026-10-10.',
'[5] やくも果樹研究所. シャインマスカット栽培 摘房 摘粒 1回目[EB/OL]. https://kaju-fruits.com/シャインマスカット　粗摘房・粗摘粒/. 访问日期2026-10-10.',
'[6] DU W. Grape inflorescence dataset[DB/OL]. Mendeley Data,2022,V1. DOI:10.17632/skwbr5k9gg.1. CC BY 4.0.',
'[7] Tom. シャインマスカットのジベレリン処理後の予備摘粒[EB/OL]. https://tomtom1ono.hatenablog.com/entry/2026/06/19/225730. 2026-06-19.',
'[8] Grapetopia. 粒間引き 摘粒 荒間引き[EB/OL]. https://grapetopia.hatenablog.com/entry/2020/06/06/222016. 2020-06-06.',
'[9] DOMINGOS S, SCAFIDI P, CARDOSO V, et al. Flower abscission in Vitis vinifera L. triggered by gibberellic acid and shade discloses differences in the underlying metabolic pathways[J]. Frontiers in Plant Science, 2015,6:457. DOI:10.3389/fpls.2015.00457.'
]
for text in refs:
 a=D.add_paragraph(text);a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.line_spacing=Pt(13)
 for run in a.runs:run.font.size=Pt(8.5)
md+=refs
D.core_properties.author='';D.core_properties.last_modified_by='';D.core_properties.title='葡萄可见果粒检测与尺度筛选算法原型设计';D.core_properties.subject='生物生产机器人设计课程论文';D.core_properties.comments=''
import tempfile,shutil
with tempfile.TemporaryDirectory() as temporary:
 staged=Path(temporary)/'paper.docx';D.save(staged);shutil.copyfile(staged,O/'课程设计论文.docx')
(O/'论文正文.md').write_text('\n\n'.join(md),encoding='utf-8');print('saved characters',len('\n'.join(md)),'figures',len(D.inline_shapes))
