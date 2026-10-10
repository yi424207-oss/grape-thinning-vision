"""Course paper built from saved actual evaluation records; anonymous cover."""
from pathlib import Path
import json,csv
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
R=Path(__file__).resolve().parents[1];O=R/'docs';O.mkdir(exist_ok=True)
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.82);sec.left_margin=sec.right_margin=Inches(.9)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 s=D.styles[name];s.font.name='Noto Serif CJK SC';s.font.color.rgb=RGBColor(0,0,0);s._element.rPr.rFonts.set(qn('w:eastAsia'),'Noto Serif CJK SC');s.font.size=Pt(11 if name=='Normal' else 16 if name=='Heading 1' else 12 if name in ['Heading 2','Caption'] else 22 if name=='Title' else 15)
 s.paragraph_format.space_after=Pt(7);s.paragraph_format.line_spacing=Pt(19 if name=='Normal' else 14 if name=='Caption' else 25 if name in ['Title','Subtitle'] else 22)
 s.font.italic=False
 for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
  s._element.rPr.rFonts.attrib.pop(qn('w:'+attr),None)
# Remove inherited Word title borders and theme overrides.
for style in D.styles:
 if style.type == 1:
  pp=style._element.find(qn('w:pPr'))
  if pp is not None:
   for b in list(pp.findall(qn('w:pBdr'))):pp.remove(b)
D.styles['Normal'].paragraph_format.first_line_indent=Pt(22)
D.styles['Caption'].font.size=Pt(9);D.styles['Caption'].paragraph_format.first_line_indent=Pt(0)
header=sec.header.paragraphs[0];header.text='生物生产机器人设计  课程设计论文';header.alignment=WD_ALIGN_PARAGRAPH.RIGHT;header.runs[0].font.size=Pt(9)
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
md=[]
def p(text):D.add_paragraph(text);md.append(text+'\n')
def h(text,level=1):D.add_heading(text,level);md.append('#'*level+' '+text+'\n')
def table(caption,headers,rows):
 cap=D.add_paragraph(caption,'Caption');cap.paragraph_format.keep_with_next=True;md.append(caption+'\n')
 t=D.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 for row in rows:
  cells=t.add_row().cells
  for c,v in zip(cells,row):c.text=str(v)
 for ri,row in enumerate(t.rows):
  no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
  for c in row.cells:
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcPr=c._tc.get_or_add_tcPr();b=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
   tcPr.append(b)
   if ri==0:
    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8E8E8');tcPr.append(shade)
   for para in c.paragraphs:
    para.paragraph_format.keep_with_next=(ri<len(t.rows)-1);para.paragraph_format.first_line_indent=Pt(0);para.paragraph_format.space_after=Pt(4);para.paragraph_format.space_before=Pt(4);para.paragraph_format.line_spacing=Pt(14)
    for run in para.runs:run.font.size=Pt(9)
  if ri==0:rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
 D.add_paragraph().paragraph_format.space_after=Pt(0)
 md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def fig(file,caption,width=6.3):
 a=D.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.first_line_indent=Pt(0);a.paragraph_format.keep_with_next=True;a.paragraph_format.line_spacing=1.0;a.add_run().add_picture(str(O/'figures'/file),width=Inches(width));b=D.add_paragraph(caption,'Caption');b.alignment=WD_ALIGN_PARAGRAPH.CENTER;md.append(f'![{caption}](figures/{file})\n')

